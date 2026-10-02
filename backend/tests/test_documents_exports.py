from io import BytesIO

from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.config import get_settings
from app.services.document_parser import DocumentParseError, extract_document
from app.services.exports import safe_csv
from conftest import make_meeting, processed_meeting
import pytest


def document_bytes(kind):
    if kind == "txt":
        return b"Maya: Action: Rahul to write deployment tests by 2026-10-08."
    if kind == "csv":
        return b"speaker,content\nMaya,Action: document the deployment\n"
    stream = BytesIO()
    if kind == "docx":
        document = Document()
        document.add_paragraph("Maya: Decision: We agreed to a canary deployment.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Rahul"
        table.cell(0, 1).text = "Action: add rollback tests."
        document.save(stream)
    else:
        pdf = canvas.Canvas(stream)
        pdf.drawString(50, 760, "Maya: Decision: We agreed to a canary deployment.")
        pdf.showPage()
        pdf.drawString(50, 760, "Rahul: Action: add rollback tests by 2026-10-08.")
        pdf.save()
    return stream.getvalue()


@pytest.mark.parametrize("kind", ["txt", "pdf", "docx", "csv"])
def test_real_document_upload_extraction_download_and_delete(client, auth, kind):
    m = make_meeting(client, auth, "")
    content = document_bytes(kind)
    result = client.post(
        "/api/documents/upload",
        headers=auth,
        data={"meeting_id": m["id"]},
        files={"file": ("transcript." + kind, content)},
    )
    assert result.status_code == 201, result.text
    doc = result.json()
    assert doc["processing_status"] == "completed" and doc["extracted_text"]
    detail = client.get(f"/api/meetings/{m['id']}", headers=auth).json()
    assert "[Document: transcript." in detail["transcript"]
    if kind == "pdf":
        assert "page 2" in detail["transcript"] and len(doc["pages"]) == 2
    assert client.get(f"/api/documents/{doc['id']}/download", headers=auth).content == content
    assert client.get("/api/documents", headers=auth).json()["total"] == 1
    assert client.delete(f"/api/documents/{doc['id']}", headers=auth).status_code == 204
    assert client.get(f"/api/documents/{doc['id']}", headers=auth).status_code == 404
    # Imported text is deliberately retained; UI explains this at deletion.
    assert (
        client.get(f"/api/meetings/{m['id']}", headers=auth).json()["transcript"]
        == detail["transcript"]
    )


def test_rejected_files_and_persisted_errors(client, auth):
    assert (
        client.post(
            "/api/documents/upload", headers=auth, files={"file": ("evil.exe", b"not a transcript")}
        ).status_code
        == 415
    )
    result = client.post(
        "/api/documents/upload", headers=auth, files={"file": ("spoofed.pdf", b"not a pdf")}
    )
    assert result.status_code == 422
    stored = client.get("/api/documents", headers=auth).json()["items"][0]
    assert stored["processing_status"] == "failed" and stored["processing_error"]
    result = client.post("/api/documents/upload", headers=auth, files={"file": ("empty.txt", b"")})
    assert result.status_code == 422


def test_file_size_and_safe_filename(client, auth):
    settings = get_settings()
    old = settings.max_upload_size
    settings.max_upload_size = 1024
    try:
        result = client.post(
            "/api/documents/upload", headers=auth, files={"file": ("large.txt", b"A" * 1100)}
        )
        assert result.status_code == 413
    finally:
        settings.max_upload_size = old
    result = client.post(
        "/api/documents/upload",
        headers=auth,
        files={"file": ("../../unsafe.txt", b"Maya: A genuine transcript line.")},
    )
    assert result.status_code == 201 and result.json()["filename"] == "unsafe.txt"
    assert "storage_path" not in result.json()


def test_standalone_attachment_invalidates_old_analysis(client, auth):
    m = processed_meeting(client, auth)
    result = client.post(
        "/api/documents/upload",
        headers=auth,
        files={"file": ("additional.txt", b"Rahul: Action: write the release checklist.")},
    )
    assert result.status_code == 201
    doc = result.json()
    assert (
        client.post(
            f"/api/documents/{doc['id']}/attach", headers=auth, json={"meeting_id": m["id"]}
        ).status_code
        == 200
    )
    result = client.get(f"/api/meetings/{m['id']}", headers=auth).json()
    assert result["summary"] is None and "release checklist" in result["transcript"]


@pytest.mark.parametrize("kind", ["pdf", "docx", "json", "markdown", "csv"])
def test_all_export_formats_have_real_content(client, auth, kind):
    m = processed_meeting(client, auth)
    result = client.get(f"/api/meetings/{m['id']}/export/{kind}", headers=auth)
    assert result.status_code == 200 and "attachment;" in result.headers["Content-Disposition"]
    if kind == "pdf":
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(result.content)).pages)
        assert "Platform sprint" in text and "Evidence" in text and "Rahul" in text
    elif kind == "docx":
        doc = Document(BytesIO(result.content))
        text = " ".join(p.text for p in doc.paragraphs)
        assert "Platform sprint" in text and "Grounding" in text and "Evidence" in text
    elif kind == "json":
        assert result.json()["action_items"] and result.json()["grounding_report"]
    elif kind == "markdown":
        assert "## Executive summary" in result.text and "> Evidence" in result.text
    else:
        assert "task,owner,due_date" in result.text and "Rahul" in result.text


def test_csv_formula_injection_guard():
    assert safe_csv("=HYPERLINK('bad')").startswith("'")
    assert safe_csv(" +cmd").startswith("'")
    assert safe_csv("2026-10-01") == "2026-10-01"


def test_parser_preserves_word_order_and_rejects_invalid_encoding():
    text, _ = extract_document("sample.docx", document_bytes("docx"))
    assert text.index("Decision") < text.index("Rahul")
    with pytest.raises(DocumentParseError):
        extract_document("binary.txt", b"\xff\xfe")
    with pytest.raises(DocumentParseError):
        extract_document("archive.docx", b"PK corrupt")


def test_csv_speaker_columns_do_not_invent_row_number_owners():
    from app.services.local_extractor import extract_local
    from app.services.preprocessing import speaker_statistics

    text, sections = extract_document(
        "notes.csv", b"speaker,content\nRahul,I'll write the runbook by 2026-10-08.\n"
    )
    result = extract_local(text)
    assert result.action_items[0].owner == "Rahul"
    assert set(speaker_statistics(text)) == {"Rahul"}
    assert sections[0]["page"] is None and sections[0]["section"] == 1
    generic, _ = extract_document("notes.csv", b"task,owner\nWrite a runbook,Rahul\n")
    assert not any(name.startswith("Row") for name in speaker_statistics(generic))


def test_raw_transcript_is_preserved_not_stripped_by_request_model(client, auth):
    text = "\n\nMaya: Original whitespace.  \n"
    m = make_meeting(client, auth, text)
    assert m["transcript"] == text
