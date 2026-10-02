"""Real JSON/Markdown/CSV/PDF/DOCX exports, including provenance and evidence."""

import csv
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


MIME_TYPES = {
    "json": "application/json",
    "markdown": "text/markdown",
    "csv": "text/csv",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _evidence_lines(item: dict) -> list[str]:
    return [
        f"Evidence · lines {e['line_start']}–{e['line_end']}"
        + (
            f" · {e['filename']}"
            + (f", page {e['page']}" if e.get("page") else f", section {e.get('section', 1)}")
            if e.get("filename")
            else ""
        )
        + f": {e['quote']}"
        for e in item.get("evidence", [])
    ]


def sections(data: dict) -> list[tuple[str, list[dict]]]:
    summary = data.get("summary") or {}
    groups = [
        ("Three-line summary", summary.get("three_line_summary", [])),
        ("Executive summary", summary.get("executive_summary", [])),
        ("Detailed summary", summary.get("detailed_summary", [])),
        ("Key points", summary.get("key_points", [])),
        ("Decisions", data["decisions"]),
        ("Action items", data["action_items"]),
        ("Risks", data["risks"]),
        ("Unanswered questions / answers", data["questions"]),
        ("Topics", [{**t, "text": t["name"]} for t in data["topics"]]),
    ]
    for key, title in (
        ("follow_ups", "Follow-up items"),
        ("commitments", "Commitments"),
        ("dependencies", "Dependencies"),
        ("projects", "Mentioned projects"),
        ("people", "Mentioned people"),
        ("important_dates", "Important dates"),
    ):
        groups.append((title, summary.get(key, [])))
    return groups


def item_text(item: dict) -> str:
    text = item.get("text", item.get("task", ""))
    if "task" in item:
        text += f"\nOwner: {item['owner']} · Due: {item['due_date_text']} · Priority: {item['priority']} · Status: {item['status']}"
        text += f"\nProvenance: {item['source_type']}" + (
            "; manually edited (original proposal retained)" if item.get("manually_edited") else ""
        )
        if item.get("notes"):
            text += f"\nNotes: {item['notes']}"
        if item.get("dependencies"):
            text += "\nDependencies: " + "; ".join(
                d.get("text", "") if isinstance(d, dict) else d for d in item["dependencies"]
            )
    if "decision_type" in item:
        text = f"[{item['decision_type']}] {text}"
    if "category" in item:
        text += f"\n{item['category']} · {item['severity']} · {item['status']}"
    if "answer" in item:
        text += f"\nStatus: {item['status']}"
        if item["answer"]:
            text += f"\nHuman-provided answer: {item['answer']}"
    support = item.get("grounding_status", "review")
    text += f"\nGrounding: {support}"
    if support not in {"user_provided", "flagged"}:
        text += f" · support score {item.get('confidence', 0):.0%} (heuristic, not probability)"
    return text


def markdown_export(data: dict) -> str:
    lines = [
        f"# {data['title']}",
        "",
        f"**Date:** {data['date']} · **Type:** {data['meeting_type']} · **Project:** {data['project'] or 'Not specified'}",
        f"**Participants (provided metadata):** {', '.join(data['participants']) or 'Not provided'}",
        "",
        f"**Analysis provider:** {data['analysis_provider'] or 'Not processed'}",
        "",
    ]
    for title, items in sections(data):
        lines += [f"## {title}", ""]
        if not items:
            lines += ["No supported items extracted.", ""]
        for item in items:
            lines += ["- " + item_text(item).replace("\n", "\n  "), ""]
            lines += ["> " + ev.replace("\n", "\n> ") for ev in _evidence_lines(item)]
            lines.append("")
    tone = (data.get("summary") or {}).get("sentiment")
    if tone:
        lines += [
            "## Tone overview",
            f"{tone['label']}: {tone['description']}",
            *_evidence_lines(tone),
            "",
        ]
    report = data.get("grounding_report") or {}
    lines += [
        "## Grounding check",
        f"Verified: {report.get('verified', 0)} · Review: {report.get('review', 0)} · Withheld: {report.get('flagged', 0)} · Metadata corrections: {report.get('metadata_corrections', 0)}",
        report.get("confidence_method", "No analysis available."),
        "",
    ]
    for entry in report.get("flagged_claims", []) + report.get("corrections", []):
        lines.append(f"- **Flagged / NOT FACT:** {entry.get('text', '')} — {entry['reason']}")
    lines += ["", "## Original transcript", ""]
    lines.extend("> " + line for line in data["transcript"].splitlines())
    return "\n".join(lines)


def safe_csv(value) -> str:
    result = "" if value is None else str(value)
    # Prevent formula execution when users open an exported CSV in spreadsheets.
    if result.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + result
    return result


def csv_export(data: dict) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(
        [
            "task",
            "owner",
            "due_date",
            "literal_deadline",
            "priority",
            "status",
            "dependencies",
            "notes",
            "source",
            "grounding",
            "confidence",
            "evidence",
        ]
    )
    for item in data["action_items"]:
        writer.writerow(
            [
                safe_csv(value)
                for value in (
                    item["task"],
                    item["owner"],
                    item["due_date"],
                    item["due_date_text"],
                    item["priority"],
                    item["status"],
                    "; ".join(
                        d.get("text", "") if isinstance(d, dict) else d
                        for d in item["dependencies"]
                    ),
                    item["notes"],
                    item["source_type"] + (" / edited" if item.get("manually_edited") else ""),
                    item["grounding_status"],
                    item["confidence"],
                    " | ".join(_evidence_lines(item)),
                )
            ]
        )
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def pdf_export(data: dict) -> bytes:
    stream = BytesIO()
    assets = Path(__file__).resolve().parents[1] / "assets"
    regular = assets / "DejaVuSans.ttf"
    if regular.exists() and "GroundedSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("GroundedSans", str(regular)))
        pdfmetrics.registerFont(TTFont("GroundedSans-Bold", str(assets / "DejaVuSans-Bold.ttf")))
    font = "GroundedSans" if regular.exists() else "Helvetica"
    bold = "GroundedSans-Bold" if regular.exists() else "Helvetica-Bold"
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
        style.leading = 15
    styles["Title"].fontName = bold
    styles["Heading2"].fontName = bold
    styles["Heading2"].textColor = colors.HexColor("#176c50")
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 14
    styles["Italic"].fontSize = 8
    styles["Italic"].textColor = colors.HexColor("#59636c")
    elements = []

    def para(text, style="BodyText"):
        return Paragraph(escape(text).replace("\n", "<br/>"), styles[style])

    elements += [
        para("GROUNDED / MEETING INTELLIGENCE", "Heading2"),
        para(data["title"], "Title"),
        para(f"{data['date']} · {data['meeting_type']} · {data['project'] or 'No project'}"),
        para("Participants: " + (", ".join(data["participants"]) or "Not provided")),
        para("Provider: " + (data["analysis_provider"] or "Not processed")),
        Spacer(1, 0.2 * inch),
    ]
    for title, items in sections(data):
        elements.append(para(title, "Heading2"))
        if not items:
            elements.append(para("No supported items extracted."))
        for item in items:
            elements += [para(item_text(item)), Spacer(1, 5)]
            for ev in _evidence_lines(item):
                elements.append(para(ev, "Italic"))
            elements.append(Spacer(1, 10))
    report = data.get("grounding_report") or {}
    elements += [
        para("Grounding check", "Heading2"),
        para(
            f"{report.get('verified', 0)} verified · {report.get('review', 0)} need review · {report.get('flagged', 0)} withheld · {report.get('metadata_corrections', 0)} corrections"
        ),
        para(report.get("confidence_method", "No analysis available.")),
    ]
    for entry in report.get("flagged_claims", []) + report.get("corrections", []):
        elements.append(
            para("Flagged / NOT FACT: " + entry.get("text", "") + " — " + entry["reason"])
        )

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#778077"))
        canvas.drawString(42, 28, "Grounded · source-backed meeting intelligence")
        canvas.drawRightString(doc.pagesize[0] - 42, 28, str(doc.page))

    SimpleDocTemplate(
        stream,
        rightMargin=42,
        leftMargin=42,
        topMargin=36,
        bottomMargin=45,
        title=data["title"],
        author="Grounded",
    ).build(elements, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def docx_export(data: dict) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(data["title"], 0)
    doc.add_paragraph(
        f"{data['date']} | {data['meeting_type']} | {data['project'] or 'No project'}"
    )
    doc.add_paragraph("Participants: " + (", ".join(data["participants"]) or "Not provided"))
    doc.add_paragraph("Analysis provider: " + (data["analysis_provider"] or "Not processed"))
    for title, items in sections(data):
        doc.add_heading(title, 1)
        if not items:
            doc.add_paragraph("No supported items extracted.")
        for item in items:
            doc.add_paragraph(item_text(item))
            for ev in _evidence_lines(item):
                paragraph = doc.add_paragraph(ev)
                paragraph.paragraph_format.left_indent = Inches(0.2)
                for run in paragraph.runs:
                    run.italic = True
                    run.font.size = Pt(9)
    report = data.get("grounding_report") or {}
    doc.add_heading("Grounding check", 1)
    doc.add_paragraph(
        f"Verified: {report.get('verified', 0)} | Review: {report.get('review', 0)} | Withheld: {report.get('flagged', 0)} | Corrections: {report.get('metadata_corrections', 0)}"
    )
    doc.add_paragraph(report.get("confidence_method", "No analysis available."))
    for entry in report.get("flagged_claims", []) + report.get("corrections", []):
        doc.add_paragraph("Flagged / NOT FACT: " + entry.get("text", "") + " — " + entry["reason"])
    doc.add_heading("Original transcript", 1)
    for line in data["transcript"].splitlines():
        doc.add_paragraph(line)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def export_meeting(data: dict, format: str) -> tuple[bytes, str, str]:
    if format == "json":
        content = json.dumps(data, ensure_ascii=False, indent=2).encode()
    elif format == "markdown":
        content = markdown_export(data).encode()
    elif format == "csv":
        content = csv_export(data)
    elif format == "pdf":
        content = pdf_export(data)
    else:
        content = docx_export(data)
    extension = "md" if format == "markdown" else format
    filename = re.sub(r"[^a-zA-Z0-9_-]+", "-", data["title"]).strip("-")[:80] or "meeting"
    return content, MIME_TYPES[format], f"{filename}.{extension}"
